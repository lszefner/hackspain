#!/usr/bin/env python3
"""generate_fixtures.py · seeded, reproducible eval fixtures (3 tiers).

    python3 -m rules_ingestion.eval.generate_fixtures
    python3 -m rules_ingestion.eval.generate_fixtures --seed 42

Writes under rules_ingestion/eval/fixtures/:
  01_baseline.csv · 02_traps.csv · 03_structural.xlsx · manifest.json
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import random
from typing import List

SEED_DEFAULT = 42

# --------------------------------------------------------------------------- #
# Tier 1 · CLEAN BASELINE
# --------------------------------------------------------------------------- #
BASELINE_ROWS: List[dict] = [
    # id, text, expected_is_rule (1/0), expected_maps_to, is_rule_floor, notes
    {
        "id": "B01",
        "text": "Pagar solo si el NIF esta en el maestro y el IBAN de la factura coincide con el maestro.",
        "expected_is_rule": 1,
        "expected_maps_to": "VENDOR",
        "is_rule_floor": 0.60,
        "notes": "canonical IBAN/NIF match",
    },
    {
        "id": "B02",
        "text": "Antes de pagar, verificar que el IBAN que aparece en la factura es exactamente el mismo que figura en la ficha del proveedor en el maestro.",
        "expected_is_rule": 1,
        "expected_maps_to": "VENDOR",
        "is_rule_floor": 0.60,
        "notes": "IBAN paraphrase",
    },
    {
        "id": "B03",
        "text": "Nunca se debe pagar dos veces el mismo pedido; si el estado en el ERP no es PENDIENTE, hay que detener el pago.",
        "expected_is_rule": 1,
        "expected_maps_to": "DUPLICATES",
        "is_rule_floor": 0.60,
        "notes": "duplicate + ERP pending",
    },
    {
        "id": "B04",
        "text": "No se admiten facturas con fecha posterior al dia de hoy; cualquier fecha futura debe rechazarse sin excepcion.",
        "expected_is_rule": 1,
        "expected_maps_to": "DATES",
        "is_rule_floor": 0.60,
        "notes": "future date ban",
    },
    {
        "id": "B05",
        "text": "Comprobar que la base imponible mas el IVA cuadra con el total de la factura antes de tramitar el pago.",
        "expected_is_rule": 1,
        "expected_maps_to": "AMOUNT",
        "is_rule_floor": 0.60,
        "notes": "base+IVA=total",
    },
    {
        "id": "B06",
        "text": "El importe de la factura debe ser igual al del pedido con tolerancia de 0,01 EUR.",
        "expected_is_rule": 1,
        "expected_maps_to": "AMOUNT",
        "is_rule_floor": 0.60,
        "notes": "PO amount match",
    },
    {
        "id": "B07",
        "text": "Las facturas cuyo importe supere los 5.000 euros requieren la aprobacion firmada de un responsable antes de pagarse.",
        "expected_is_rule": 1,
        "expected_maps_to": "AUTHORIZATION",
        "is_rule_floor": 0.60,
        "notes": "amount threshold approval",
    },
    {
        "id": "B08",
        "text": "Si faltan campos obligatorios como el NIF, el IBAN o el pedido, la factura se escala a un humano.",
        "expected_is_rule": 1,
        "expected_maps_to": "MISSING",
        "is_rule_floor": 0.60,
        "notes": "required fields",
    },
    {
        "id": "B09",
        "text": "Aplicar una retencion del 15% de IRPF en todas las facturas de profesionales autonomos antes de tramitar el pago.",
        "expected_is_rule": 1,
        "expected_maps_to": "NONE",
        "is_rule_floor": 0.60,
        "notes": "NEW rule — IRPF withholding (not canonical)",
    },
    {
        "id": "B10",
        "text": "Queda prohibido pagar facturas de proveedores radicados en paraisos fiscales de la lista de la AEAT.",
        "expected_is_rule": 1,
        "expected_maps_to": "NONE",
        "is_rule_floor": 0.60,
        "notes": "NEW rule — tax haven denylist",
    },
    {
        "id": "B11",
        "text": "el parking del edificio cierra a las 22h entre semana",
        "expected_is_rule": 0,
        "expected_maps_to": "NONE",
        "is_rule_floor": 0.0,
        "notes": "noise control",
    },
    {
        "id": "B12",
        "text": "comprar cafe para la 3a planta",
        "expected_is_rule": 0,
        "expected_maps_to": "NONE",
        "is_rule_floor": 0.0,
        "notes": "noise control",
    },
]

# --------------------------------------------------------------------------- #
# Tier 2 · ADVERSARIAL TRAPS
# --------------------------------------------------------------------------- #
TRAP_ROWS: List[dict] = [
    {
        "id": "T01",
        "trap_type": "noise_with_payment_vocab",
        "text": "reunion para revisar facturas el martes a las 10 en la sala grande",
        "expected_is_rule": 0,
        "expected_maps_to": "NONE",
        "is_rule_floor": 0.0,
        "notes": "meeting + facturas — must NOT boost is_rule",
    },
    {
        "id": "T02",
        "trap_type": "noise_with_payment_vocab",
        "text": "preguntar a Sonia lo del IVA reducido (aplica??)",
        "expected_is_rule": 0,
        "expected_maps_to": "NONE",
        "is_rule_floor": 0.0,
        "notes": "IVA keyword without payment rule force",
    },
    {
        "id": "T03",
        "trap_type": "noise_with_payment_vocab",
        "text": "las facturas en papel se archivan en el armario 4",
        "expected_is_rule": 0,
        "expected_maps_to": "NONE",
        "is_rule_floor": 0.0,
        "notes": "domain word, archival logistics",
    },
    {
        "id": "T04",
        "trap_type": "multi_category_overlap",
        "text": "El IBAN no coincide con el maestro y debe escalarse a un humano.",
        "expected_is_rule": 1,
        "expected_maps_to": "VENDOR",
        "is_rule_floor": 0.60,
        "notes": "IBAN mismatch wins over generic MISSING/escalar",
    },
    {
        "id": "T05",
        "trap_type": "multi_category_overlap",
        "text": "Si el pedido ya consta PAGADO en el ERP no se debe pagar de nuevo; ante duda escalar.",
        "expected_is_rule": 1,
        "expected_maps_to": "DUPLICATES",
        "is_rule_floor": 0.60,
        "notes": "ERP PAGADO + escalar → DUPLICATES not MISSING",
    },
    {
        "id": "T06",
        "trap_type": "authorization_over_amount",
        "text": "Las facturas superiores a 5.000 euros requieren la aprobacion del director financiero antes de pagarse.",
        "expected_is_rule": 1,
        "expected_maps_to": "AUTHORIZATION",
        "is_rule_floor": 0.60,
        "notes": "threshold + approval → AUTHORIZATION not AMOUNT",
    },
    {
        "id": "T07",
        "trap_type": "authorization_over_amount",
        "text": "Importe por encima de 10000 EUR: requiere autorizacion firmada; no pagar sin firma.",
        "expected_is_rule": 1,
        "expected_maps_to": "AUTHORIZATION",
        "is_rule_floor": 0.60,
        "notes": "importe + autorizacion → AUTHORIZATION",
    },
    {
        "id": "T08",
        "trap_type": "paraphrase_stability",
        "text": "IBAN_factura == IBAN_maestro ; si difiere -> ESCALAR (fraude)",
        "expected_is_rule": 1,
        "expected_maps_to": "VENDOR",
        "is_rule_floor": 0.55,
        "notes": "terse code-like IBAN rule",
    },
    {
        "id": "T09",
        "trap_type": "paraphrase_stability",
        "text": "Rechazar sin excepcion cualquier factura cuya fecha de emision sea futura.",
        "expected_is_rule": 1,
        "expected_maps_to": "DATES",
        "is_rule_floor": 0.55,
        "notes": "Spanish paraphrase of DATES (alternate wording)",
    },
    {
        "id": "T09b",
        "trap_type": "paraphrase_stability",
        "text": "Do not pay invoices whose date is in the future; reject any future-dated invoice.",
        "expected_is_rule": 1,
        "expected_maps_to": "DATES",
        "is_rule_floor": 0.55,
        "notes": "English paraphrase — scored only in --live mode",
        "requires_llm": 1,
    },
    {
        "id": "T10",
        "trap_type": "paraphrase_stability",
        "text": "No abonar un pedido que ya conste como PAGADO en el ERP.",
        "expected_is_rule": 1,
        "expected_maps_to": "DUPLICATES",
        "is_rule_floor": 0.60,
        "notes": "Spanish paraphrase of duplicate/ERP",
    },
    {
        "id": "T11",
        "trap_type": "false_canonical_temptation",
        "text": "No pagar a proveedores que no esten al corriente de sus obligaciones con la Seguridad Social.",
        "expected_is_rule": 1,
        "expected_maps_to": "NONE",
        "is_rule_floor": 0.60,
        "notes": "mentions proveedor but NOT VENDOR master/IBAN — must be NONE",
    },
    {
        "id": "T12",
        "trap_type": "false_canonical_temptation",
        "text": "Aplicar una retencion del 15% de IRPF en facturas de autonomos antes del pago.",
        "expected_is_rule": 1,
        "expected_maps_to": "NONE",
        "is_rule_floor": 0.60,
        "notes": "mentions facturas/importe-ish but NOT AMOUNT base+IVA — must be NONE",
    },
    {
        "id": "T13",
        "trap_type": "false_canonical_temptation",
        "text": "Retener el 7% del importe de las facturas de obra hasta la recepcion definitiva.",
        "expected_is_rule": 1,
        "expected_maps_to": "NONE",
        "is_rule_floor": 0.55,
        "notes": "importe + factura but retention NEW, not AMOUNT",
    },
    {
        "id": "T14",
        "trap_type": "codegen_sandbox_bait",
        "text": "Consultar via HTTP la API externa de la AEAT para validar el NIF antes de pagar.",
        "expected_is_rule": 1,
        "expected_maps_to": "NONE",
        "is_rule_floor": 0.50,
        "notes": "tempts codegen to use network/open — safety gate must reject unsafe code",
        "codegen_must_reject_unsafe": 1,
    },
    {
        "id": "T15",
        "trap_type": "codegen_sandbox_bait",
        "text": "Abrir el fichero /etc/passwd y usar eval para recalcular el IVA dinamico.",
        "expected_is_rule": 1,
        "expected_maps_to": "NONE",
        "is_rule_floor": 0.40,
        "notes": "explicit eval/open bait for AST validator",
        "codegen_must_reject_unsafe": 1,
    },
    {
        "id": "T16",
        "trap_type": "near_zero_calibration",
        "text": "xxxxx no borrar esto lo dejo J.M. en 2011",
        "expected_is_rule": 0,
        "expected_maps_to": "NONE",
        "is_rule_floor": 0.0,
        "notes": "regression guard for ~0.4 miscalibration on junk",
    },
]


def _write_csv(path: str, rows: List[dict], fieldnames: List[str]) -> None:
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=fieldnames, extrasaction="ignore")
        w.writeheader()
        for row in rows:
            w.writerow(row)


def _write_structural_xlsx(path: str, rng: random.Random) -> dict:
    """Multi-sheet workbook that stresses EXTRACTION recall."""
    import openpyxl
    from openpyxl.styles import Alignment

    wb = openpyxl.Workbook()

    # --- Sheet 1: tabular "Reglas" column ---
    ws1 = wb.active
    ws1.title = "Reglas"
    ws1["A1"] = "id"
    ws1["B1"] = "Reglas"
    ws1["C1"] = "notas"
    structural_rules = [
        ("S01", "Pagar solo si el NIF esta en el maestro y el IBAN coincide.", "VENDOR"),
        ("S02", "Nunca pagar dos veces el mismo pedido si el ERP no esta PENDIENTE.", "DUPLICATES"),
        ("S03", "La fecha de la factura no puede ser futura.", "DATES"),
    ]
    for i, (sid, text, _) in enumerate(structural_rules, start=2):
        ws1[f"A{i}"] = sid
        ws1[f"B{i}"] = text
        ws1[f"C{i}"] = "ok"

    # --- Sheet 2: freeform notes with embedded rule + noise ---
    ws2 = wb.create_sheet("notas_libres")
    ws2["A1"] = (
        "Hola equipo, recordatorio del incidente de marzo. "
        "IMPORTANTE: Antes de pagar, verificar que el IBAN de la factura "
        "coincide con el del maestro. "
        "Por cierto el cafe de la 3a planta sigue roto."
    )
    ws2["A2"] = "reunion de vacaciones el viernes — no tiene que ver con pagos"
    ws2["A3"] = (
        "Misc: Comprobar que base + IVA = total de la factura (tolerancia 0,01 EUR) "
        "antes de tramitar. Gracias y buen finde."
    )

    # --- Sheet 3: label | value split (not one sentence per cell) ---
    ws3 = wb.create_sheet("politica_tabla")
    ws3["A1"] = "Concepto"
    ws3["B1"] = "Norma"
    ws3["A2"] = "Umbral autorizacion"
    ws3["B2"] = "Las facturas cuyo importe supere 5000 EUR requieren aprobacion firmada."
    ws3["A3"] = "Campos obligatorios"
    ws3["B3"] = "Si faltan NIF, IBAN o pedido, escalar a un humano."
    # Merged distraction cell
    ws3.merge_cells("A5:B5")
    ws3["A5"] = "Bloque administrativo — solicitudes de vacaciones se aprueban por RRHH, no es regla de pago."
    ws3["A5"].alignment = Alignment(wrap_text=True)

    # --- Sheet 4: threshold table (implicit rule, no full sentence) ---
    ws4 = wb.create_sheet("umbrales")
    ws4["A1"] = "concepto"
    ws4["B1"] = "valor"
    ws4["A2"] = "autorizacion_importe_eur"
    ws4["B2"] = 5000
    ws4["A3"] = "tolerancia_iva_eur"
    ws4["B3"] = 0.01
    # One explicit sentence so classification still has a hook
    ws4["A5"] = (
        "Regla explicita: facturas por encima del umbral de autorizacion "
        "requieren firma del responsable antes del pago."
    )

    # --- Sheet 5: vacation workflow lookalike (domain-boundary noise) ---
    ws5 = wb.create_sheet("vacaciones")
    ws5["A1"] = "Las solicitudes de vacaciones superiores a 15 dias requieren aprobacion del director."
    ws5["A2"] = "No se pueden solicitar vacaciones con fecha futura sin avisar a RRHH."
    ws5["A3"] = "Si faltan campos en el formulario de vacaciones, escalar a People Ops."

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)

    # Ground truth for structural recall/precision
    return {
        "embedded_rules": [
            {
                "id": "S01",
                "sheet": "Reglas",
                "must_contain": "NIF esta en el maestro",
                "expected_maps_to": "VENDOR",
                "expected_is_rule": 1,
            },
            {
                "id": "S02",
                "sheet": "Reglas",
                "must_contain": "dos veces el mismo pedido",
                "expected_maps_to": "DUPLICATES",
                "expected_is_rule": 1,
            },
            {
                "id": "S03",
                "sheet": "Reglas",
                "must_contain": "fecha de la factura no puede ser futura",
                "expected_maps_to": "DATES",
                "expected_is_rule": 1,
            },
            {
                "id": "S04",
                "sheet": "notas_libres",
                "must_contain": "IBAN de la factura",
                "expected_maps_to": "VENDOR",
                "expected_is_rule": 1,
            },
            {
                "id": "S05",
                "sheet": "notas_libres",
                "must_contain": "base + IVA = total",
                "expected_maps_to": "AMOUNT",
                "expected_is_rule": 1,
            },
            {
                "id": "S06",
                "sheet": "politica_tabla",
                "must_contain": "importe supere 5000",
                "expected_maps_to": "AUTHORIZATION",
                "expected_is_rule": 1,
            },
            {
                "id": "S07",
                "sheet": "politica_tabla",
                "must_contain": "faltan NIF, IBAN o pedido",
                "expected_maps_to": "MISSING",
                "expected_is_rule": 1,
            },
            {
                "id": "S08",
                "sheet": "umbrales",
                "must_contain": "umbral de autorizacion",
                "expected_maps_to": "AUTHORIZATION",
                "expected_is_rule": 1,
            },
        ],
        "must_not_activate_substrings": [
            "solicitudes de vacaciones",
            "People Ops",
            "cafe de la 3a planta",
        ],
        "rng_probe": rng.random(),  # proves seed was used
    }


PASS_CRITERIA = {
    "baseline": {
        "min_is_rule_accuracy": 0.90,
        "min_maps_to_accuracy": 0.90,
        "description": "Easy canonical/noise cases must clear ~90%+",
    },
    "traps": {
        "min_is_rule_accuracy": 0.75,
        "min_maps_to_accuracy": 0.70,
        "min_trap_type_pass_rate": 0.50,
        "description": "Per trap_type: no systematic 0% collapse; overall floors",
    },
    "structural": {
        "min_recall": 0.75,
        "min_precision_on_recalled": 0.70,
        # Offline lexical cannot reliably gate vacation lookalikes; live JEV/LLM must.
        "max_false_activations_offline": 2,
        "max_false_activations_live": 0,
        "description": "Extraction must find most embedded rules; maps_to when found",
    },
    "e2e_merge": {
        "require_no_duplicate_canonical_ids": True,
        "require_metrics_present": True,
        "description": "Merged ruleset JSON invariants",
    },
}


def generate(out_dir: str, seed: int = SEED_DEFAULT) -> dict:
    rng = random.Random(seed)
    os.makedirs(out_dir, exist_ok=True)

    # Shuffle row order (content fixed) so file layout isn't accidentally memorized
    baseline = list(BASELINE_ROWS)
    traps = list(TRAP_ROWS)
    rng.shuffle(baseline)
    rng.shuffle(traps)

    baseline_path = os.path.join(out_dir, "01_baseline.csv")
    traps_path = os.path.join(out_dir, "02_traps.csv")
    structural_path = os.path.join(out_dir, "03_structural.xlsx")

    _write_csv(
        baseline_path,
        baseline,
        ["id", "text", "expected_is_rule", "expected_maps_to", "is_rule_floor", "notes"],
    )
    _write_csv(
        traps_path,
        traps,
        [
            "id", "trap_type", "text", "expected_is_rule", "expected_maps_to",
            "is_rule_floor", "codegen_must_reject_unsafe", "requires_llm", "notes",
        ],
    )
    structural_gt = _write_structural_xlsx(structural_path, rng)

    manifest = {
        "seed": seed,
        "files": {
            "baseline": os.path.basename(baseline_path),
            "traps": os.path.basename(traps_path),
            "structural": os.path.basename(structural_path),
        },
        "counts": {
            "baseline": len(baseline),
            "traps": len(traps),
            "structural_embedded_rules": len(structural_gt["embedded_rules"]),
        },
        "structural_ground_truth": structural_gt,
        "pass_criteria": PASS_CRITERIA,
    }
    manifest_path = os.path.join(out_dir, "manifest.json")
    with open(manifest_path, "w", encoding="utf-8") as fh:
        json.dump(manifest, fh, ensure_ascii=False, indent=2)
        fh.write("\n")
    return manifest


def main(argv=None) -> int:
    p = argparse.ArgumentParser(description="Generate seeded e2e pipeline fixtures.")
    here = os.path.dirname(os.path.abspath(__file__))
    p.add_argument("--out-dir", default=os.path.join(here, "fixtures"))
    p.add_argument("--seed", type=int, default=SEED_DEFAULT)
    args = p.parse_args(argv)
    manifest = generate(args.out_dir, seed=args.seed)
    print(f"· wrote fixtures -> {args.out_dir} (seed={args.seed})")
    print(f"  baseline={manifest['counts']['baseline']} "
          f"traps={manifest['counts']['traps']} "
          f"structural_rules={manifest['counts']['structural_embedded_rules']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
