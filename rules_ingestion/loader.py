"""loader.py · read any workbook per sources.yaml -> clean lookups + schema.

The loader is deliberately dumb about *meaning*: it applies the sources.yaml
contract (which sheets, which columns, which normalizer per column) and emits:

  * lookups     : {"proveedores": {id: row}, "pedidos": {pedido: row}}  (normalized)
  * norma_lines : [{"text","cell","row"}]  raw rule candidates for the classifier
  * schema      : per-sheet discovery (expected vs found vs missing vs extra)
  * warnings    : anything a human should know (dupes, dropped junk, blanks)

Discovering schema (not assuming it) is what lets Alberto drop a new sheet or a
renamed column and see it in the report instead of a silent wrong answer.
"""
from __future__ import annotations

import csv
import io
import os
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Dict, List, Optional

import yaml

from . import normalize as N

_NORMALIZERS = {
    "text": N.norm_text,
    "nif": N.norm_nif,
    "iban": N.norm_iban,
    "importe": N.norm_importe,
    "fecha": N.norm_fecha,
}


@dataclass
class LoadResult:
    lookups: Dict[str, Dict[str, dict]] = field(default_factory=dict)
    norma_lines: List[dict] = field(default_factory=list)
    schema: Dict[str, dict] = field(default_factory=dict)
    warnings: List[str] = field(default_factory=list)
    ruleset_version: str = "v3"
    source_records: Dict[str, List[dict]] = field(default_factory=dict)
    source_provenance: Dict[str, List[dict]] = field(default_factory=dict)
    workbook_bytes: bytes = b""
    source_config_bytes: bytes = b""


def _cell_safe(value: Any) -> Any:
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bool) or value is None \
            or isinstance(value, (str, int, float)):
        return value
    return str(value)


def load_config(path: str) -> dict:
    with open(path, "rb") as fh:
        return yaml.safe_load(fh.read())


def _apply_normalizers(row: dict, norm_map: Dict[str, str]) -> dict:
    out = dict(row)
    for col, kind in (norm_map or {}).items():
        fn = _NORMALIZERS.get(kind)
        if fn is not None and col in out:
            out[col] = fn(out[col])
    return out


def _read_sheet_rows(ws, header_row: int) -> "tuple[list, list]":
    """Return (headers, data_rows) as lists; headers stripped strings."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [N.norm_text(h) or "" for h in rows[header_row - 1]]
    data = rows[header_row:]
    return headers, data


def _load_tabular(ws, sheet_cfg: dict, sheet_key: str, result: LoadResult,
                  base_dir: str = ".") -> None:
    headers, data = _read_sheet_rows(ws, sheet_cfg.get("header_row") or 1)
    col_map: Dict[str, str] = sheet_cfg["columns"]           # canonical -> source header
    norm_map: Dict[str, str] = sheet_cfg.get("normalize", {})
    # Campos que la hoja no trae pero la fuente afirma igualmente. La Caja no
    # tiene columna de moneda y sus PDFs a menudo tampoco la imprimen; el
    # maestro de pedidos SI es autoridad sobre order.currency. Declararlo aqui
    # lo deja en el registro, con el sha de este fichero en la traza, en vez de
    # como un default escondido en el codigo.
    const_map: Dict[str, Any] = sheet_cfg.get("constants", {}) or {}
    key_col = sheet_cfg["key"]

    # --- schema discovery ---
    found = [h for h in headers if h]
    expected = list(col_map.values())
    missing = [h for h in expected if h not in found]
    extra = [h for h in found if h not in expected]
    result.schema[sheet_key] = {
        "sheet": sheet_cfg["sheet"],
        "expected_columns": expected,
        "found_columns": found,
        "missing_columns": missing,
        "extra_columns": extra,
        "data_rows": len(data),
    }
    if missing:
        result.warnings.append(
            f"[{sheet_key}] missing declared columns {missing} in sheet "
            f"'{sheet_cfg['sheet']}' (found {found})"
        )

    # source header -> position index
    index_of = {h: i for i, h in enumerate(headers)}

    table: Dict[str, dict] = {}
    dupes = 0
    blanks = 0
    captured: List[dict] = []
    provenance: List[dict] = []
    header_row = sheet_cfg.get("header_row") or 1
    from openpyxl.utils import get_column_letter
    for offset, raw in enumerate(data):
        row_number = header_row + 1 + offset
        record: Dict[str, Any] = {}
        cells: Dict[str, str] = {}
        raw_values: Dict[str, Any] = {}
        for canonical, source_header in col_map.items():
            idx = index_of.get(source_header)
            if idx is not None:
                cells[canonical] = f"{get_column_letter(idx + 1)}{row_number}"
            value = raw[idx] if (idx is not None and idx < len(raw)) else None
            raw_values[canonical] = _cell_safe(value)
            record[canonical] = value
        record = _apply_normalizers(record, norm_map)
        for canonical, constant in const_map.items():
            record.setdefault(canonical, constant)
        captured.append({k: _cell_safe(v) for k, v in record.items()})
        provenance.append({"sheet": sheet_cfg["sheet"], "row": row_number,
                           "cells": cells, "raw": raw_values})
        key_val = record.get(key_col)
        if key_val in (None, ""):
            blanks += 1
            continue
        key_val = str(key_val)
        if key_val in table:
            dupes += 1          # last-write-wins; handles duplicate P007 row
        table[key_val] = record

    if dupes:
        result.warnings.append(
            f"[{sheet_key}] {dupes} duplicate key(s) collapsed on '{key_col}' "
            f"(last-write-wins)"
        )
    if blanks:
        result.warnings.append(f"[{sheet_key}] {blanks} row(s) with blank key skipped")
    # Fuentes incrementales: los CSV que Alberto deja caer junto al Excel
    # (proveedores_nuevos.csv, pedidos_nuevos.csv del lote del sabado). Se
    # funden en la MISMA tabla y por la MISMA clave, asi que el resto del
    # motor no se entera de que un proveedor vino de una hoja o de un CSV.
    # Cada CSV declara su propio mapa de columnas porque los encabezados no
    # tienen por que coincidir con los de la hoja.
    for extra in sheet_cfg.get("merge_csv", []) or []:
        csv_path = extra["path"]
        if not os.path.isabs(csv_path):
            csv_path = os.path.join(base_dir, csv_path)
        if not os.path.isfile(csv_path):
            result.warnings.append(
                f"[{sheet_key}] fuente incremental declarada y ausente: {csv_path}")
            continue
        nuevos = 0
        with open(csv_path, "r", encoding="utf-8-sig", newline="") as fh:
            for number, raw_row in enumerate(csv.DictReader(fh), start=2):
                record = {canonical: raw_row.get(header)
                          for canonical, header in extra["columns"].items()}
                record = _apply_normalizers(record, norm_map)
                for canonical, constant in const_map.items():
                    record.setdefault(canonical, constant)
                key_val = record.get(key_col)
                if key_val in (None, ""):
                    continue
                captured.append({k: _cell_safe(v) for k, v in record.items()})
                provenance.append({"sheet": os.path.basename(csv_path),
                                   "row": number, "cells": {}, "raw": dict(raw_row)})
                table[str(key_val)] = record
                nuevos += 1
        result.warnings.append(
            f"[{sheet_key}] {nuevos} registro(s) de {os.path.basename(csv_path)}")

    result.lookups[sheet_key] = table
    result.source_records[sheet_key] = captured
    result.source_provenance[sheet_key] = provenance


def _load_norma(ws, norma_cfg: dict, result: LoadResult) -> None:
    col_letter = (norma_cfg.get("text_column") or "A").upper()
    col_idx = 0
    for ch in col_letter:                       # A->0, B->1, ... (single/multi)
        col_idx = col_idx * 26 + (ord(ch) - ord("A") + 1)
    col_idx -= 1
    skip_res = [re.compile(p) for p in norma_cfg.get("skip_if_matches", [])]
    result.ruleset_version = norma_cfg.get("ruleset_version", result.ruleset_version)
    declared = norma_cfg.get("declared_rules", True)  # norm sheet = declared rules

    lines: List[dict] = []
    for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
        val = row[col_idx] if col_idx < len(row) else None
        text = N.norm_text(val)
        if not text:
            continue
        if any(rx.search(text) for rx in skip_res):
            continue
        lines.append({"text": text, "cell": f"{col_letter}{r}", "row": r,
                      "declared": declared, "source_sheet": norma_cfg["sheet"]})
    result.norma_lines = lines
    result.schema["norma"] = {
        "sheet": norma_cfg["sheet"],
        "ruleset_version": result.ruleset_version,
        "candidate_lines": len(lines),
    }


def load(config_path: str, base_dir: Optional[str] = None) -> LoadResult:
    """Load everything declared in sources.yaml. base_dir defaults to config dir."""
    import openpyxl

    with open(config_path, "rb") as fh:
        config_bytes = fh.read()
    cfg = yaml.safe_load(config_bytes)
    base_dir = base_dir or os.path.dirname(os.path.abspath(config_path))
    wb_path = os.path.join(base_dir, cfg["workbook"]["path"])
    with open(wb_path, "rb") as fh:
        workbook_bytes = fh.read()
    wb = openpyxl.load_workbook(io.BytesIO(workbook_bytes),
                              read_only=True, data_only=True)

    result = LoadResult()
    result.workbook_bytes = workbook_bytes
    result.source_config_bytes = config_bytes

    # record which sheets we deliberately dropped (auditability)
    declared = {c["sheet"] for c in cfg["workbook"]["sheets"].values()}
    ignored = set(cfg["workbook"].get("ignore_sheets", []))
    for name in wb.sheetnames:
        if name not in declared and name not in ignored:
            result.warnings.append(
                f"[schema] sheet '{name}' present in workbook but neither "
                f"declared nor ignored in sources.yaml -> treated as junk"
            )

    sheets_cfg = cfg["workbook"]["sheets"]
    for sheet_key, sheet_cfg in sheets_cfg.items():
        name = sheet_cfg["sheet"]
        if name not in wb.sheetnames:
            result.warnings.append(f"[{sheet_key}] declared sheet '{name}' not found")
            continue
        ws = wb[name]
        if sheet_key == "norma":
            _load_norma(ws, sheet_cfg, result)
        else:
            _load_tabular(ws, sheet_cfg, sheet_key, result, base_dir)

    wb.close()
    return result


def is_text_candidate(value: object, min_len: int = 15, min_words: int = 3) -> bool:
    """Cheap pre-filter: does this cell look like a sentence (a rule candidate)?

    Keeps JEV cost bounded — we only classify prose, not codes/numbers/dates.
    """
    if not isinstance(value, str):
        return False
    s = value.strip()
    if len(s) < min_len or len(s.split()) < min_words:
        return False
    letters = sum(c.isalpha() for c in s)
    if letters < len(s) * 0.4:            # mostly digits/punctuation -> skip
        return False
    return True


def scan_candidates(folder: str, min_len: int = 15) -> List[dict]:
    """Walk a folder of CSV/XLSX and emit sentence-like rule candidates.

    Each candidate: {file, sheet, cell, text, declared=False}. `declared=False`
    means the classifier runs the FULL Noul(is_rule)+Choice(maps_to) — this is
    DISCOVERY: finding rules hidden among random text and across sheets.
    """
    import openpyxl
    from openpyxl.utils import get_column_letter

    out: List[dict] = []
    for name in sorted(os.listdir(folder)):
        path = os.path.join(folder, name)
        if not os.path.isfile(path):
            continue
        ext = os.path.splitext(name)[1].lower()
        if ext in (".xlsx", ".xlsm"):
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            for ws in wb.worksheets:
                for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
                    for c, val in enumerate(row, start=1):
                        text = N.norm_text(val)
                        if text and is_text_candidate(text, min_len):
                            out.append({"file": name, "sheet": ws.title,
                                        "cell": f"{get_column_letter(c)}{r}",
                                        "text": text, "declared": False})
            wb.close()
        elif ext == ".csv":
            with open(path, "r", encoding="utf-8-sig", newline="") as fh:
                for r, row in enumerate(csv.reader(fh), start=1):
                    for c, val in enumerate(row, start=1):
                        text = N.norm_text(val)
                        if text and is_text_candidate(text, min_len):
                            out.append({"file": name, "sheet": "",
                                        "cell": f"{get_column_letter(c)}{r}",
                                        "text": text, "declared": False})
    return out


def load_csv_lookup(path: str, key_col: str, norm_map: Optional[Dict[str, str]] = None
                    ) -> Dict[str, dict]:
    """Load a CSV export (e.g. Saturday's lote2) into a normalized lookup.

    Kept separate + simple so a new CSV connector is config-shaped, not a rewrite.
    """
    table: Dict[str, dict] = {}
    with open(path, "r", encoding="utf-8-sig", newline="") as fh:
        for record in csv.DictReader(fh):
            record = _apply_normalizers(dict(record), norm_map or {})
            key_val = record.get(key_col)
            if key_val not in (None, ""):
                table[str(key_val)] = record
    return table
