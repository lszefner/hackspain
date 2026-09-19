from __future__ import annotations

import hashlib
import json
import urllib.request
from pathlib import Path

import pytest
import yaml
from test_decision_context import (
    CAPTURED,
    EVAL_DATE,
    build,
    make_outcome,
    make_ruleset,
    make_snapshots,
)

from rules_ingestion import loader
from rules_ingestion.decision_context import evaluate_context
from rules_ingestion.decision_storage import create_decision_store
from webui.master_data import (
    ErpClient,
    capture_erp_snapshot,
    capture_master_snapshots,
)
from webui.results_store import ResultsStore
from webui.run_revision import review_outcome


def _workbook(tmp_path: Path, extra_rows=None, drop_sheet=None,
              dup_header=None):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Proveedores"
    ws.append(["ID", "Razon Social", "NIF", "IBAN", "Ciudad", "Condiciones"])
    if dup_header:
        ws.cell(row=1, column=7, value=dup_header)
    ws.append(["P001", "Proveedor Uno", "B12345678",
               "ES91 2100 0418 4502 0005 1332", "Madrid", "30 dias"])
    ws.append(["P001", "Proveedor Uno Dup", "B12345678",
               "ES91 2100 0418 4502 0005 1332", "Madrid", "30 dias"])
    for row in extra_rows or []:
        ws.append(row)
    ws2 = wb.create_sheet("Pedidos_2026")
    ws2.append(["Pedido", "ProveedorID", "NIF", "Importe_Total", "Estado",
                "Fecha_Pedido"])
    ws2.append(["PO-1", "P001", "B12345678", "121.00", "ABIERTO",
                "2026-08-01"])
    path = tmp_path / "master.xlsx"
    tmp_path.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    cfg = {
        "workbook": {
            "path": path.name,
            "sheets": {
                "proveedores": {
                    "sheet": "Proveedores", "header_row": 1, "key": "id",
                    "columns": {"id": "ID", "razon_social": "Razon Social",
                                "nif": "NIF", "iban": "IBAN",
                                "ciudad": "Ciudad",
                                "condiciones": "Condiciones"},
                    "normalize": {"nif": "nif", "iban": "iban",
                                  "razon_social": "text",
                                  "condiciones": "text"},
                },
                "pedidos": {
                    "sheet": "Pedidos_2026", "header_row": 1, "key": "pedido",
                    "columns": {"pedido": "Pedido",
                                "proveedor_id": "ProveedorID", "nif": "NIF",
                                "importe_total": "Importe_Total",
                                "estado_excel": "Estado",
                                "fecha_pedido": "Fecha_Pedido"},
                    "normalize": {"nif": "nif", "importe_total": "importe",
                                  "fecha_pedido": "fecha"},
                },
            },
        }
    }
    if drop_sheet:
        del cfg["workbook"]["sheets"][drop_sheet]
    yaml_path = tmp_path / "sources.yaml"
    yaml_path.write_text(yaml.safe_dump(cfg), encoding="utf-8")
    return yaml_path, path


def test_loader_captures_records_and_provenance(tmp_path):
    yaml_path, wb_path = _workbook(tmp_path)
    result = loader.load(str(yaml_path))
    suppliers = result.source_records["proveedores"]
    assert len(suppliers) == 2
    assert all(r["id"] == "P001" for r in suppliers)
    assert list(result.lookups["proveedores"]) == ["P001"]
    assert result.lookups["proveedores"]["P001"]["razon_social"] \
        == "Proveedor Uno Dup"
    prov = result.source_provenance["proveedores"]
    assert prov[0]["row"] == 2 and prov[0]["cells"]["nif"] == "C2"
    assert prov[0]["raw"]["nif"] == "B12345678"
    assert result.workbook_bytes == wb_path.read_bytes()
    assert result.source_config_bytes == yaml_path.read_bytes()


def test_capture_master_snapshots_scopes_and_status(tmp_path):
    yaml_path, wb_path = _workbook(tmp_path)
    snaps = capture_master_snapshots(str(yaml_path), captured_at=CAPTURED)
    assert set(snaps) == {"workbook", "source_mapping", "suppliers", "orders"}
    assert snaps["workbook"].payload == wb_path.read_bytes()
    assert snaps["source_mapping"].payload == yaml_path.read_bytes()
    suppliers = snaps["suppliers"]
    assert suppliers.kind == "supplier_master"
    assert "supplier.active" not in suppliers.authoritative_for
    assert suppliers.availability == "available"
    assert len(suppliers.payload["records"]) == 2
    assert suppliers.payload["workbook_sha256"] == \
        hashlib.sha256(wb_path.read_bytes()).hexdigest()
    orders = snaps["orders"]
    assert "order.currency" not in orders.authoritative_for
    assert orders.payload["records"][0]["pedido"] == "PO-1"


def test_capture_master_snapshots_missing_sheet(tmp_path):
    yaml_path, _ = _workbook(tmp_path, drop_sheet="pedidos")
    snaps = capture_master_snapshots(str(yaml_path), captured_at=CAPTURED)
    assert snaps["orders"].availability == "unavailable"
    assert snaps["suppliers"].availability == "available"


def test_capture_master_snapshots_missing_configured_sheet(tmp_path):
    yaml_path, _ = _workbook(tmp_path)
    cfg = yaml.safe_load(yaml_path.read_text())
    cfg["workbook"]["sheets"]["pedidos"]["sheet"] = "NoExiste"
    yaml_path.write_text(yaml.safe_dump(cfg), encoding="utf-8")
    snaps = capture_master_snapshots(str(yaml_path), captured_at=CAPTURED)
    assert snaps["orders"].availability == "unavailable"
    assert snaps["suppliers"].availability == "available"


def test_capture_master_snapshots_broken_source(tmp_path):
    bad = tmp_path / "missing.yaml"
    bad.write_text("workbook: {path: 'nope.xlsx', sheets: {}}")
    snaps = capture_master_snapshots(str(bad), captured_at=CAPTURED)
    assert snaps["suppliers"].availability == "unavailable"


def test_capture_master_snapshots_identity_problems(tmp_path):
    yaml_path, _ = _workbook(tmp_path)
    cfg = yaml.safe_load(yaml_path.read_text())
    del cfg["workbook"]["sheets"]["proveedores"]["columns"]["id"]
    yaml_path.write_text(yaml.safe_dump(cfg), encoding="utf-8")
    snaps = capture_master_snapshots(str(yaml_path), captured_at=CAPTURED)
    assert snaps["suppliers"].availability == "partial"

    yaml_path, _ = _workbook(tmp_path / "a")
    cfg = yaml.safe_load(yaml_path.read_text())
    cfg["workbook"]["sheets"]["proveedores"]["columns"]["id"] = "NOEXISTE"
    yaml_path.write_text(yaml.safe_dump(cfg), encoding="utf-8")
    snaps = capture_master_snapshots(str(yaml_path), captured_at=CAPTURED)
    assert snaps["suppliers"].availability == "partial"

    yaml_path, _ = _workbook(tmp_path / "b", dup_header="NIF")
    snaps = capture_master_snapshots(str(yaml_path), captured_at=CAPTURED)
    assert snaps["suppliers"].availability == "partial"


class _FakeResponse:
    def __init__(self, body: bytes):
        self._body = body

    def __enter__(self):
        return self

    def __exit__(self, *args):
        return False

    def read(self):
        return self._body


def _erp_xml(rows, pagina=1, paginas=1, total=None, por_pagina=None):
    filas = "".join(
        f"<asiento><pedido>{p}</pedido><estado>{e}</estado></asiento>"
        for p, e in rows)
    total = len(rows) if total is None else total
    por_pagina = max(len(rows), 1) if por_pagina is None else por_pagina
    return (
        f"<asientos>{filas}<meta><total>{total}</total>"
        f"<paginas>{paginas}</paginas><pagina>{pagina}</pagina>"
        f"<por_pagina>{por_pagina}</por_pagina></meta></asientos>")


def _mock_erp(monkeypatch, pages: dict,
              login_xml=b"<login><token>SECRETTOKEN</token></login>",
              login_xmls=None):
    calls = []

    def fake_urlopen(req, timeout=None):
        calls.append(req)
        url = req.full_url if hasattr(req, "full_url") else req
        if "/erp/login" in url:
            if login_xmls:
                return _FakeResponse(
                    login_xmls.pop(0) if login_xmls else login_xml)
            return _FakeResponse(login_xml)
        import re as _re
        page = int(_re.search(r"pagina=(\d+)", url).group(1))
        result = pages.get(page)
        if isinstance(result, Exception):
            raise result
        if result is None:
            raise urllib.error.URLError("boom")
        return _FakeResponse(result)

    monkeypatch.setattr(urllib.request, "urlopen", fake_urlopen)
    monkeypatch.setattr("time.sleep", lambda *_: None)
    return calls


def test_erp_snapshot_complete_two_pages(monkeypatch):
    page1 = _erp_xml([("PO-1", "PENDIENTE"), ("PO-1", "PAGADA")],
                     pagina=1, paginas=2, total=3, por_pagina=2).encode()
    page2 = _erp_xml([("PO-2", "PAGADA")],
                     pagina=2, paginas=2, total=3, por_pagina=2).encode()
    _mock_erp(monkeypatch, {1: page1, 2: page2})
    client = ErpClient()
    result = client.snapshot()
    assert result["complete"] is True
    assert len(result["pages"]) == 2
    estados = [r["estado"] for r in result["records"] if r["pedido"] == "PO-1"]
    assert estados == ["PENDIENTE", "PAGADA"]
    assert all(r["page"] and r["row"] for r in result["records"])
    snap = capture_erp_snapshot(client=client, captured_at=CAPTURED)
    assert snap.availability == "available"
    assert snap.authoritative_for == ("order.payment_status",)
    text = json.dumps(snap.payload)
    assert "SECRETTOKEN" not in text and "FACTURAS2009" not in text


def test_erp_snapshot_page_failure_preserves_progress(monkeypatch):
    page1 = _erp_xml([("PO-1", "PENDIENTE")],
                     pagina=1, paginas=3, total=3, por_pagina=1).encode()
    _mock_erp(monkeypatch, {1: page1, 2: urllib.error.URLError("down")})
    result = ErpClient().snapshot()
    assert result["complete"] is False
    assert len(result["pages"]) == 1
    assert result["error_code"] == "connection_error"
    snap = capture_erp_snapshot(client=ErpClient(), captured_at=CAPTURED)
    assert snap.availability == "partial"


def test_erp_snapshot_truncation_and_malformed(monkeypatch):
    page1 = _erp_xml([("PO-1", "PENDIENTE")],
                     pagina=1, paginas=5, total=5, por_pagina=1).encode()
    page2 = _erp_xml([("PO-2", "PENDIENTE")],
                     pagina=2, paginas=5, total=5, por_pagina=1).encode()
    _mock_erp(monkeypatch, {1: page1, 2: page2})
    result = ErpClient().snapshot(max_pages=2)
    assert result["complete"] is False
    assert result["error_code"] == "truncated_max_pages"

    bad = b"<asientos><asiento><pedido>PO-1</pedido></asiento>" \
        b"<meta><total>1</total><paginas>1</paginas><pagina>1</pagina>" \
        b"<por_pagina>5</por_pagina></meta></asientos>"
    _mock_erp(monkeypatch, {1: bad})
    result = ErpClient().snapshot()
    assert result["complete"] is False
    assert result["error_code"] == "malformed_record"
    assert len(result["pages"]) == 1 and not result["records"]

    empty = b"<asientos><meta><total>5</total><paginas>2</paginas>" \
        b"<pagina>1</pagina><por_pagina>5</por_pagina></meta></asientos>"
    _mock_erp(monkeypatch, {1: empty})
    result = ErpClient().snapshot()
    assert result["complete"] is False
    assert result["error_code"] == "empty_page"


def test_erp_snapshot_metadata_integrity(monkeypatch):
    wrong_page = _erp_xml([("PO-1", "PENDIENTE")], pagina=9, paginas=1,
                          total=1, por_pagina=5).encode()
    _mock_erp(monkeypatch, {1: wrong_page})
    assert ErpClient().snapshot()["error_code"] == "page_mismatch"

    page1 = _erp_xml([("PO-1", "PENDIENTE")], pagina=1, paginas=2,
                     total=2, por_pagina=1).encode()
    page2 = _erp_xml([("PO-2", "PENDIENTE")], pagina=2, paginas=3,
                     total=2, por_pagina=1).encode()
    _mock_erp(monkeypatch, {1: page1, 2: page2})
    result = ErpClient().snapshot()
    assert result["complete"] is False
    assert result["error_code"] == "pagination_changed"

    short = _erp_xml([("PO-2", "PENDIENTE")], pagina=2, paginas=2,
                     total=3, por_pagina=1).encode()
    page1_3 = _erp_xml([("PO-1", "PENDIENTE")], pagina=1, paginas=2,
                       total=3, por_pagina=1).encode()
    _mock_erp(monkeypatch, {1: page1_3, 2: short})
    result = ErpClient().snapshot()
    assert result["complete"] is False
    assert result["error_code"] == "count_mismatch"


def test_erp_snapshot_sensitive_response_never_stored(monkeypatch):
    leaked = _erp_xml([("PO-1", "PENDIENTE")], pagina=1, paginas=1,
                      total=1, por_pagina=5)
    leaked = leaked.replace("</asientos>",
                            "<nota>SECRETTOKEN</nota></asientos>").encode()
    _mock_erp(monkeypatch, {1: leaked})
    client = ErpClient()
    result = client.snapshot()
    assert result["complete"] is False
    assert result["error_code"] == "sensitive_response"
    assert result["pages"] == []


def test_erp_snapshot_missing_token_and_max_pages(monkeypatch):
    _mock_erp(monkeypatch, {}, login_xml=b"<login><token></token></login>")
    result = ErpClient().snapshot()
    assert result["complete"] is False
    assert result["error_code"] == "invalid_response"
    assert ErpClient().snapshot(max_pages=0)["error_code"] \
        == "invalid_max_pages"


def test_erp_snapshot_no_response_unavailable(monkeypatch):
    _mock_erp(monkeypatch, {1: urllib.error.URLError("down")},
              login_xml=b"garbage")
    client = ErpClient()
    snap = capture_erp_snapshot(client=client, captured_at=CAPTURED)
    assert snap.availability == "unavailable"
    assert "FACTURAS2009" not in snap.scope


def test_results_store_context_columns_and_history(tmp_path):
    store = ResultsStore(tmp_path / "r.db")
    bundle = build()
    evaluation = evaluate_context(bundle)
    store.set_hecha("f-clean", decision=evaluation["decision"],
                    raw_invoice={"invoice_number": "INV-1", "total": "121.00",
                                 "vendor_id": "P042", "currency": "EUR",
                                 "date": "2026-09-01"},
                    checks=evaluation["checks"], context=bundle.context,
                    receipt={"context_id": bundle.context["context_id"]})
    row = store.get("f-clean")
    assert row["decision_context"] and row["context_receipt"]
    snap = store.processed_history_snapshot(captured_at=CAPTURED)
    assert snap.kind == "history"
    assert snap.payload["kind"] == "processed"
    assert snap.payload["complete"] is False
    assert snap.availability == "partial"
    record = snap.payload["records"][0]
    assert record["supplier_id"] == "P042"
    assert record["currency"] == "EUR"
    assert record["context_id"] == bundle.context["context_id"]

    store.set_procesando("f-clean")
    row = store.get("f-clean")
    assert row["decision"] is None and row["checks"] is None
    assert row["decision_context"] is not None
    assert row["raw_invoice"] is not None
    snap = store.processed_history_snapshot(captured_at=CAPTURED)
    assert any(r["file_id"] == "f-clean" for r in snap.payload["records"])
    store.set_error("f-clean", "boom")
    row = store.get("f-clean")
    assert row["decision"] is None
    assert row["decision_context"] is not None
    snap = store.processed_history_snapshot(captured_at=CAPTURED)
    assert any(r["file_id"] == "f-clean" for r in snap.payload["records"])


def test_review_outcome_local_store(tmp_path):
    store = create_decision_store("local", local_root=tmp_path)
    bundle, evaluation, receipt = review_outcome(
        make_outcome(), snapshots=make_snapshots(),
        ruleset=make_ruleset(), evaluation_date=EVAL_DATE,
        captured_at=CAPTURED, decision_store=store)
    assert evaluation["decision"] == "PAGAR"
    assert receipt["context_id"] == bundle.context["context_id"]
    assert store.load(receipt["context_id"]).context == bundle.context
    store.close()


def test_review_outcome_persistence_failure_propagates(tmp_path):
    class FailingStore:
        def save(self, bundle):
            raise OSError("synthetic persistence failure")

    with pytest.raises(IOError):
        review_outcome(
            make_outcome(), snapshots=make_snapshots(),
            ruleset=make_ruleset(), evaluation_date=EVAL_DATE,
            captured_at=CAPTURED, decision_store=FailingStore())


def test_run_revision_guards_before_providers(monkeypatch):
    import webui.run_revision as rr

    assert not hasattr(rr, "load_dotenv")
    monkeypatch.delenv("HELMCODE_API_KEY", raising=False)
    monkeypatch.delenv("REVISION_EVALUATION_DATE", raising=False)
    with pytest.raises(ValueError):
        rr.revisar_lote_sync(["x.pdf"], store=None)
    with pytest.raises(ValueError):
        rr.revisar_lote_sync(["x.pdf"], store=None,
                             evaluation_date="not-a-date")
    monkeypatch.setenv("REVISION_BACKEND", "bogus")
    with pytest.raises(ValueError):
        rr.revisar_lote_sync(["x.pdf"], store=None,
                             evaluation_date=EVAL_DATE)
    monkeypatch.delenv("REVISION_BACKEND")
    with pytest.raises(ValueError):
        rr.revisar_lote_sync(["x.pdf"], store=None,
                             evaluation_date=EVAL_DATE,
                             ruleset_path="/nonexistent/rules.json")


def test_map_invoice_conservative(tmp_path):
    from test_decision_context import make_invoice

    from webui.map_invoice import to_raw_invoice

    invoice = make_invoice()
    raw = to_raw_invoice("f-clean", invoice)
    assert raw["iva"] == "21.00"
    assert raw["currency"] == "EUR"
    invoice["taxes"] = [{"label": "IVA 21%", "amount": "21.00"},
                        {"label": "IRPF -15%", "amount": "-15.00"}]
    assert to_raw_invoice("f-clean", invoice)["iva"] is None
    invoice["taxes"] = [{"label": "otros", "amount": "5.00"}]
    assert to_raw_invoice("f-clean", invoice)["iva"] is None
    invoice["taxes"] = []
    invoice["lines"] = [{"amount": "10.00"}, {"amount": None}]
    raw = to_raw_invoice("f-clean", invoice)
    assert raw["iva"] is None
    assert raw["line_items"] == ["10.00", None]


def test_decision_cli_roundtrip(tmp_path):
    from rules_ingestion.decision_cli import main

    outcome_path = tmp_path / "outcome.json"
    outcome_path.write_text(json.dumps(make_outcome()))
    ruleset_path = tmp_path / "rules.json"
    ruleset_path.write_bytes(
        __import__("ingestion.storage", fromlist=["canonical_json_bytes"])
        .canonical_json_bytes(make_ruleset()))
    snapshots_path = tmp_path / "snapshots.json"
    specs = {}
    for source_id, snap in make_snapshots().items():
        specs[source_id] = {
            "kind": snap.kind, "payload": snap.payload,
            "authoritative_for": list(snap.authoritative_for),
            "scope": snap.scope, "availability": snap.availability,
            "asserted_by": snap.asserted_by,
        }
    snapshots_path.write_text(json.dumps(specs))
    code = main([
        "--outcome", str(outcome_path),
        "--ruleset", str(ruleset_path),
        "--snapshots", str(snapshots_path),
        "--evaluation-date", EVAL_DATE,
        "--backend", "local",
        "--local-root", str(tmp_path / "store"),
        "--captured-at", CAPTURED,
    ])
    assert code == 0
    receipts = list((tmp_path / "store" / "contexts").glob("*.json"))
    assert len(receipts) == 1
    receipt = json.loads(receipts[0].read_text())
    assert receipt["decision"] == "PAGAR"


def test_decision_cli_failures(tmp_path):
    from rules_ingestion.decision_cli import main

    with pytest.raises(SystemExit):
        main(["--outcome", "x"])
    outcome_path = tmp_path / "outcome.json"
    outcome_path.write_text(json.dumps(make_outcome()))
    code = main([
        "--outcome", str(outcome_path),
        "--ruleset", str(tmp_path / "missing.json"),
        "--snapshots", str(tmp_path / "missing2.json"),
        "--evaluation-date", EVAL_DATE,
        "--backend", "local",
        "--local-root", str(tmp_path / "store"),
    ])
    assert code == 1


def test_map_invoice_rejects_non_finite_decimals(tmp_path):
    from test_decision_context import make_invoice

    from webui.map_invoice import to_raw_invoice

    for bad in ("NaN", "Infinity", "-Infinity", "1e3", "abc", "1.2.3"):
        invoice = make_invoice()
        invoice["taxes"] = [{"label": "IVA 21%", "amount": bad}]
        assert to_raw_invoice("f-clean", invoice)["iva"] is None, bad


class _FakePipeline:
    rows = None
    extra_rows = None
    corrupt_artifacts = False
    mismatch_invoice_file_id = False

    def __init__(self, repo, storage, contracts, config, secrets):
        self.repo = repo
        self.storage = storage
        _FakePipeline.config = config

    async def run(self, batch):
        for item in batch["manifest"]:
            file_id = item["file_id"]
            specs = (type(self).rows or {}).get(file_id, ["ok"])
            inp = self.repo.register_input(
                batch["id"], file_id, file_name=file_id)
            self.repo.ensure_job(
                batch_id=batch["id"], input_id=inp["id"],
                stage="interpret", provider="deepseek", model="test-model",
                work_key=inp["id"])
            for index, spec in enumerate(specs):
                interpreter = "deepseek" if index == 0 \
                    else f"deepseek-{index}"
                outcome = make_outcome(file_id=file_id)
                if type(self).mismatch_invoice_file_id:
                    outcome["invoice"]["file_id"] = "other.pdf"
                data = json.dumps(outcome).encode()
                ref = self.storage.put(data, "outcome")
                if type(self).corrupt_artifacts:
                    (self.storage.root / ref.object_key) \
                        .write_bytes(b"corrupted")
                art = self.repo.save_artifact(ref, payload=outcome)
                self.repo.set_result(
                    inp["id"], interpreter, art["id"],
                    "completed" if spec == "ok" else spec)
        for extra in (type(self).extra_rows or []):
            inp = self.repo.register_input(
                batch["id"], extra, file_name=extra)
            self.repo.ensure_job(
                batch_id=batch["id"], input_id=inp["id"],
                stage="interpret", provider="deepseek", model="test-model",
                work_key=inp["id"])
            self.repo.set_result(inp["id"], "deepseek", None, "failed")
        return {"ok": True}

    def load_artifact(self, artifact_id):
        artifact = self.repo.get_artifact(artifact_id)
        data = self.storage.get(artifact["object_key"])
        assert hashlib.sha256(data).hexdigest() == artifact["sha256"]
        return json.loads(data)


def _revision_fixture(tmp_path, monkeypatch, file_ids=("f-run.pdf",)):
    import webui.run_revision as rr

    facturas = tmp_path / "facturas"
    facturas.mkdir(exist_ok=True)
    for file_id in file_ids:
        (facturas / file_id).write_bytes(b"%PDF-fake")
    data_dir = tmp_path / "data"
    monkeypatch.setattr(rr, "FACTURAS_DIR", facturas)
    monkeypatch.setattr(rr, "DATA_DIR", data_dir)
    monkeypatch.setenv("HELMCODE_API_KEY", "test-key")
    monkeypatch.delenv("REVISION_BACKEND", raising=False)
    monkeypatch.setattr(ErpClient, "snapshot", lambda self, **kw: {
        "records": [], "pages": [], "complete": False,
        "error_code": "connection_error"})
    monkeypatch.setattr(rr, "Pipeline", _FakePipeline)
    yaml_path, _ = _workbook(tmp_path / "src")
    ruleset_path = tmp_path / "rules.json"
    ruleset_path.write_text(json.dumps(make_ruleset()))
    return rr, {"evaluation_date": EVAL_DATE,
                "ruleset_path": str(ruleset_path),
                "sources_yaml": str(yaml_path)}


def test_revisar_lote_full_path_local(tmp_path, monkeypatch):
    rr, kwargs = _revision_fixture(tmp_path, monkeypatch)
    store = ResultsStore(tmp_path / "rev.db")
    result = rr.revisar_lote_sync(["f-run.pdf"], store=store, **kwargs)
    assert result["revision_counts"] == {"stored": 1, "failed": 0}
    assert _FakePipeline.config["schemas"]
    row = store.get("f-run.pdf")
    assert row["estado"] == "hecha"
    assert row["decision"] in ("PAGAR", "ESCALAR", "NO_PAGAR")
    assert row["decision_context"] and row["context_receipt"]
    contexts_dir = tmp_path / "data" / "decision-contexts" / "contexts"
    assert len(list(contexts_dir.glob("*.json"))) == 1


def test_revisar_lote_persistence_failure_retains_evidence(
        tmp_path, monkeypatch):
    rr, kwargs = _revision_fixture(tmp_path, monkeypatch)
    store = ResultsStore(tmp_path / "rev.db")
    store.set_hecha("f-run.pdf", decision="PAGAR",
                    raw_invoice={"invoice_number": "OLD"},
                    checks=[], context={"context_id": "dc_old"},
                    receipt={"backend": "local"})

    class FailingStore:
        repository = None

        def save(self, bundle):
            raise OSError("synthetic persistence failure")

        def close(self):
            pass

    monkeypatch.setattr(rr, "create_decision_store",
                        lambda *a, **kw: FailingStore())
    result = rr.revisar_lote_sync(["f-run.pdf"], store=store, **kwargs)
    assert result["revision_counts"] == {"stored": 0, "failed": 1}
    row = store.get("f-run.pdf")
    assert row["estado"] == "error"
    assert row["decision"] is None
    assert row["decision_context"] is not None
    snap = store.processed_history_snapshot(captured_at=CAPTURED)
    assert any(r["file_id"] == "f-run.pdf" for r in snap.payload["records"])


def test_revisar_lote_file_id_validation_before_work(
        tmp_path, monkeypatch):
    rr, kwargs = _revision_fixture(tmp_path, monkeypatch)

    def _no_store(*a, **kw):
        raise AssertionError("store must not be created")

    monkeypatch.setattr(rr, "create_decision_store", _no_store)
    store = ResultsStore(tmp_path / "rev.db")
    for bad in ("../evil.pdf", "a/b.pdf", "a\\b.pdf", "/abs.pdf",
                "noext", "x.PDF.exe"):
        with pytest.raises(ValueError):
            rr.revisar_lote_sync([bad], store=store, **kwargs)
    with pytest.raises(ValueError):
        rr.revisar_lote_sync([], store=store, **kwargs)
    with pytest.raises(ValueError):
        rr.revisar_lote_sync(["f-run.pdf", "f-run.pdf"], store=store,
                             **kwargs)


def test_revisar_lote_malformed_ruleset_before_pipeline(
        tmp_path, monkeypatch):
    rr, kwargs = _revision_fixture(tmp_path, monkeypatch)
    (tmp_path / "bad-rules.json").write_bytes(b"{")
    kwargs["ruleset_path"] = str(tmp_path / "bad-rules.json")
    monkeypatch.setattr(rr, "settings",
                        lambda *a, **kw: (_ for _ in ()).throw(
                            AssertionError("settings must not run")))
    monkeypatch.setattr(rr, "create_decision_store",
                        lambda *a, **kw: (_ for _ in ()).throw(
                            AssertionError("store must not run")))
    from rules_ingestion.decision_context import ContextError
    with pytest.raises((ValueError, ContextError)):
        rr.revisar_lote_sync(["f-run.pdf"],
                             store=ResultsStore(tmp_path / "rev.db"),
                             **kwargs)


def test_server_detail_escapes_and_labels(tmp_path, monkeypatch):
    import sys
    import types

    import webui.run_revision as rr

    (tmp_path / "srv-data").mkdir()
    monkeypatch.setattr(rr, "DATA_DIR", tmp_path / "srv-data")
    sys.modules.pop("webui.server", None)
    from webui import server

    facturas = tmp_path / "facturas"
    facturas.mkdir()
    (facturas / "f1.pdf").write_bytes(b"%PDF")
    monkeypatch.setattr(server, "FACTURAS_DIR", facturas)
    store = ResultsStore(tmp_path / "rev.db")
    store.set_hecha("f1.pdf", decision="PAGAR",
                    raw_invoice={"invoice_number": "<script>x</script>"},
                    checks=[{"canonical": "<b>chk</b>", "verdict": "PAGAR",
                             "reason": "<i>r</i>"}],
                    context={"context_id": "dc_x", "sources": {}},
                    receipt={"backend": "local"})
    monkeypatch.setattr(server, "STORE", store)
    captured = {}
    fake = types.SimpleNamespace(
        _responder=lambda status, body, **kw: captured.update(body=body))
    server.Handler._pagina_factura(fake, "f1.pdf")
    body = captured["body"]
    assert "<script>x</script>" not in body
    assert "&lt;script&gt;x&lt;/script&gt;" in body
    assert "recomendacion, no pago" in body

    store.set_error("f1.pdf", "boom")
    captured.clear()
    server.Handler._pagina_factura(fake, "f1.pdf")
    body = captured["body"]
    assert "no recomendacion vigente" in body
    assert ">PAGAR<" not in body


def test_decision_cli_redacts_exception_details(tmp_path, monkeypatch,
                                              capsys):
    import webui.run_revision  # noqa: F401
    from rules_ingestion import decision_cli

    outcome_path = tmp_path / "outcome.json"
    outcome_path.write_text(json.dumps(make_outcome()))
    ruleset_path = tmp_path / "rules.json"
    ruleset_path.write_text(json.dumps(make_ruleset()))
    snapshots_path = tmp_path / "snapshots.json"
    specs = {sid: {"kind": s.kind, "payload": s.payload,
                   "authoritative_for": list(s.authoritative_for)}
             for sid, s in make_snapshots().items()}
    snapshots_path.write_text(json.dumps(specs))

    def _explode(*a, **kw):
        raise RuntimeError(
            "dsn=postgres://user:SECRET_PW@host token=SECRETTOKEN")

    monkeypatch.setattr(decision_cli, "create_decision_store", _explode)
    code = decision_cli.main([
        "--outcome", str(outcome_path),
        "--ruleset", str(ruleset_path),
        "--snapshots", str(snapshots_path),
        "--evaluation-date", EVAL_DATE,
        "--backend", "local",
        "--local-root", str(tmp_path / "store"),
        "--captured-at", CAPTURED,
    ])
    assert code == 1
    err = capsys.readouterr().err
    assert "SECRET_PW" not in err and "SECRETTOKEN" not in err
    assert "decision_context_failed" in err


def test_erp_snapshot_refresh_empty_token_stops(monkeypatch):
    page1 = urllib.error.HTTPError(
        "u", 401, "unauth", None, None)
    _mock_erp(monkeypatch, {1: page1},
              login_xmls=[b"<login><token>SECRETTOKEN</token></login>",
                          b"<login><token></token></login>"])
    result = ErpClient().snapshot()
    assert result["complete"] is False
    assert result["error_code"] == "invalid_response"
    assert result["pages"] == []


def test_revisar_lote_row_accounting(tmp_path, monkeypatch):
    rr, kwargs = _revision_fixture(
        tmp_path, monkeypatch, file_ids=("f-a.pdf", "f-b.pdf"))
    monkeypatch.setattr(_FakePipeline, "rows",
                        {"f-a.pdf": ["ok", "ok"], "f-b.pdf": []})
    monkeypatch.setattr(_FakePipeline, "extra_rows", ["ghost.pdf"])
    store = ResultsStore(tmp_path / "rev.db")
    result = rr.revisar_lote_sync(["f-a.pdf", "f-b.pdf"], store=store,
                                  **kwargs)
    assert result["revision_counts"] == {"stored": 0, "failed": 2}
    assert store.get("f-a.pdf")["estado"] == "error"
    assert store.get("f-b.pdf")["estado"] == "error"
    assert store.get("f-a.pdf")["decision"] is None
    contexts_dir = tmp_path / "data" / "decision-contexts" / "contexts"
    assert not contexts_dir.exists() or not list(contexts_dir.glob("*.json"))


def test_revisar_lote_corrupt_artifact_and_mismatch(tmp_path, monkeypatch):
    rr, kwargs = _revision_fixture(tmp_path, monkeypatch)
    monkeypatch.setattr(_FakePipeline, "corrupt_artifacts", True)
    store = ResultsStore(tmp_path / "rev.db")
    result = rr.revisar_lote_sync(["f-run.pdf"], store=store, **kwargs)
    assert result["revision_counts"] == {"stored": 0, "failed": 1}
    assert store.get("f-run.pdf")["estado"] == "error"
    contexts_dir = tmp_path / "data" / "decision-contexts" / "contexts"
    assert not contexts_dir.exists() or not list(contexts_dir.glob("*.json"))

    monkeypatch.setattr(_FakePipeline, "corrupt_artifacts", False)
    monkeypatch.setattr(_FakePipeline, "mismatch_invoice_file_id", True)
    store2 = ResultsStore(tmp_path / "rev2.db")
    result = rr.revisar_lote_sync(["f-run.pdf"], store=store2, **kwargs)
    assert result["revision_counts"] == {"stored": 0, "failed": 1}
    assert store2.get("f-run.pdf")["estado"] == "error"


def test_decision_cli_checked_in_examples(tmp_path):
    from jsonschema import Draft202012Validator

    from rules_ingestion.decision_cli import main
    from rules_ingestion.decision_context import load_schema
    from rules_ingestion.decision_storage import create_decision_store

    examples = Path(__file__).resolve().parent.parent \
        / "docs" / "examples" / "decision-context"
    schema = load_schema()
    Draft202012Validator.check_schema(schema)
    local_root = tmp_path / "store"
    code = main([
        "--outcome", str(examples / "outcome.json"),
        "--ruleset", str(examples / "ruleset.json"),
        "--snapshots", str(examples / "snapshots.json"),
        "--evaluation-date", EVAL_DATE,
        "--backend", "local",
        "--local-root", str(local_root),
        "--captured-at", CAPTURED,
    ])
    assert code == 0
    receipts = list((local_root / "contexts").glob("*.json"))
    assert len(receipts) == 1
    receipt = json.loads(receipts[0].read_text())
    store = create_decision_store("local", local_root=local_root)
    loaded = store.load(receipt["context_id"])
    store.close()
    Draft202012Validator(schema).validate(loaded.context)
    assert loaded.context["preflight"]["data_status"] in ("ready", "blocked")
